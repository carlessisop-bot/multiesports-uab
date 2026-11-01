import streamlit as st
import pandas as pd
import datetime
import io
import os
import re

# ==========================================
# 1. CONFIGURACIÓ DE PÀGINA I ESTILS UAB
# ==========================================
st.set_page_config(
    page_title="MultiEsports UAB - Portal Oficial & Gestió",
    page_icon="⚽",
    layout="wide"
)

# Estils CSS Corporatius UAB
st.markdown("""
    <style>
    .main-title {
        color: #008037;
        font-family: 'Helvetica', sans-serif;
        font-weight: bold;
        font-size: 2.2rem;
        margin-bottom: 0.2rem;
    }
    .sub-title {
        color: #003366;
        font-size: 1.1rem;
        margin-bottom: 1.5rem;
    }
    .badge-status {
        background-color: #e8f5e9;
        color: #2e7d32;
        padding: 4px 8px;
        border-radius: 4px;
        font-weight: bold;
    }
    .stTable {
        border-radius: 8px;
        overflow: hidden;
    }
    </style>
""", unsafe_allow_html=True)

st.markdown("""<div class="main-title">MultiEsports UAB — SAF Servei d'Activitat Física</div>""", unsafe_allow_html=True)
st.markdown("""<div class="sub-title">Portal Públic de Resultats & Panell d'Àrbitres</div>""", unsafe_allow_html=True)

PASSWORD_ARBITRE = "UAB2026"  # Contrasenya per defecte per als àrbitres/coordinadors

# ==========================================
# 2. FUNCIONS AUXILIARS I DE SEGURETAT
# ==========================================

def sanitize_pdf_text(text):
    if text is None:
        return ""
    s = str(text)
    s = s.replace("—", "-").replace("–", "-").replace("“", '"').replace("”", '"').replace("’", "'").replace("‘", "'")
    return s.encode('latin-1', 'replace').decode('latin-1')

def safe_int(val, default=0):
    try:
        if pd.notnull(val) and str(val).strip() != "" and str(val).strip().lower() != "nan":
            return int(float(val))
    except (ValueError, TypeError):
        pass
    return default

def obtenir_fitxer_excel():
    if os.path.exists("Plantilla_MultiEsports_UAB.xlsx"):
        return "Plantilla_MultiEsports_UAB.xlsx"
    elif os.path.exists("Plantilla_MultiEsports_UAB.xlsm"):
        return "Plantilla_MultiEsports_UAB.xlsm"
    else:
        files = [f for f in os.listdir(".") if (f.endswith(".xlsm") or f.endswith(".xlsx")) and not f.startswith("~$")]
        if files:
            return files[0]
    return "Plantilla_MultiEsports_UAB.xlsx"

# ==========================================
# 3. LECTURA I ESCRIPTURA D'EXCEL
# ==========================================

def carregar_calendari():
    excel_file = obtenir_fitxer_excel()
    if os.path.exists(excel_file):
        try:
            return pd.read_excel(excel_file, sheet_name="Calendari", engine="openpyxl")
        except Exception as e:
            st.error(f"Error en llegir el full Calendari del fitxer '{excel_file}': {e}")
            return pd.DataFrame()
    else:
        st.warning(f"No s'ha trobat cap fitxer Excel (.xlsx/.xlsm) a la carpeta actual ({os.getcwd()}).")
        return pd.DataFrame()

def carregar_consolidat():
    excel_file = obtenir_fitxer_excel()
    if os.path.exists(excel_file):
        try:
            xls = pd.ExcelFile(excel_file, engine="openpyxl")
            if "Consolidat" in xls.sheet_names:
                return pd.read_excel(xls, sheet_name="Consolidat")
        except Exception:
            pass
    return pd.DataFrame()

def carregar_jugadors_equip(esport, equip):
    df_cons = carregar_consolidat()
    if df_cons.empty:
        return []
    
    # Identificar columnes
    col_comp = next((c for c in df_cons.columns if "competi" in c.lower() or "esport" in c.lower()), df_cons.columns[0])
    col_eq = next((c for c in df_cons.columns if "equip" in c.lower()), df_cons.columns[1] if len(df_cons.columns)>1 else df_cons.columns[0])
    col_nom = next((c for c in df_cons.columns if "nom" in c.lower()), df_cons.columns[3] if len(df_cons.columns)>3 else df_cons.columns[0])
    col_id = next((c for c in df_cons.columns if "idparticipant" in c.lower() or "id_participant" in c.lower() or "id" in c.lower()), None)
    
    mask = (df_cons[col_comp].astype(str).str.strip().str.upper() == str(esport).strip().upper()) & \
           (df_cons[col_eq].astype(str).str.strip().str.upper() == str(equip).strip().upper())
           
    df_sub = df_cons[mask]
    jugadors = []
    for _, r in df_sub.iterrows():
        nom = str(r[col_nom]).strip()
        id_part = str(r[col_id]).strip() if col_id and pd.notnull(r[col_id]) else ""
        if nom and nom.lower() != "nan":
            jugadors.append({"nom": nom, "id_participant": id_part})
    return jugadors

def guardar_acta_amb_assistencies(id_partit, esport, eq_loc, eq_vis, res_loc, res_vis, 
                                  tg_loc=0, tv_loc=0, tg_vis=0, tv_vis=0, 
                                  text_gols_loc="", text_gols_vis="", 
                                  assistencies_loc={}, assistencies_vis={}):
    excel_file = obtenir_fitxer_excel()
    try:
        import openpyxl
        
        wb = openpyxl.load_workbook(excel_file, keep_vba=True)
        if "Calendari" not in wb.sheetnames:
            return False, "No s'ha trobat el full 'Calendari' a l'Excel."
            
        ws_cal = wb["Calendari"]
        
        # 1. Mapatge de columnes de Calendari
        col_map = {}
        for col in range(1, ws_cal.max_column + 1):
            hdr = str(ws_cal.cell(row=1, column=col).value).strip()
            if hdr: col_map[hdr] = col
            
        col_id = col_map.get("Id", 1)
        col_rl = col_map.get("ResultatLocal", 12)
        col_rv = col_map.get("ResultatVisitant", 13)
        col_estat = col_map.get("EstatPartit", 15)
        
        col_tgl = col_map.get("TargetesGroguesLocal", 18)
        col_tvl = col_map.get("TargetesVermellesLocal", 19)
        col_tgv = col_map.get("TargetesGroguesVisitant", 20)
        col_tvv = col_map.get("TargetesVermellesVisitant", 21)
        
        col_gl = col_map.get("GolejadorsLocal", 22)
        col_gv = col_map.get("GolejadorsVisitant", 23)
        
        trobat = False
        target_id_str = str(id_partit).strip()
        
        for row in range(2, ws_cal.max_row + 1):
            val_id_raw = ws_cal.cell(row=row, column=col_id).value
            if val_id_raw is not None:
                val_id_str = str(val_id_raw).strip()
                if val_id_str == target_id_str or safe_int(val_id_raw) == safe_int(id_partit):
                    ws_cal.cell(row=row, column=col_rl, value=safe_int(res_loc))
                    ws_cal.cell(row=row, column=col_rv, value=safe_int(res_vis))
                    ws_cal.cell(row=row, column=col_estat, value="VALIDAT")
                    
                    ws_cal.cell(row=row, column=col_tgl, value=safe_int(tg_loc))
                    ws_cal.cell(row=row, column=col_tvl, value=safe_int(tv_loc))
                    ws_cal.cell(row=row, column=col_tgv, value=safe_int(tg_vis))
                    ws_cal.cell(row=row, column=col_tvv, value=safe_int(tv_vis))
                    
                    if text_gols_loc: ws_cal.cell(row=row, column=col_gl, value=str(text_gols_loc))
                    if text_gols_vis: ws_cal.cell(row=row, column=col_gv, value=str(text_gols_vis))
                    trobat = True
                    break
                    
        if not trobat:
            return False, f"No s'ha trobat el partit amb ID '{id_partit}' al Calendari."
            
        # 2. Gravació d'Assistències al full 'Assistencies'
        if "Assistencies" not in wb.sheetnames:
            ws_ass = wb.create_sheet("Assistencies")
            headers_ass = ["IdPartit", "Competicio", "Equip", "NomJugador", "IdParticipant", "Assisteix", "DataRegistre"]
            for c_idx, h in enumerate(headers_ass, 1):
                ws_ass.cell(row=1, column=c_idx, value=h)
        else:
            ws_ass = wb["Assistencies"]
            
        now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
        
        # Eliminar registres anteriors d'aquest partit si ja s'havia validat
        rows_to_delete = []
        for r in range(2, ws_ass.max_row + 1):
            if str(ws_ass.cell(row=r, column=1).value).strip() == target_id_str:
                rows_to_delete.append(r)
        for r in reversed(rows_to_delete):
            ws_ass.delete_rows(r)
            
        # Assegurar capçaleres correctes a la fila 1 de Assistencies
        headers_ass = ["IdPartit", "IdParticipant", "NomJugador", "Competició", "Equip", "Assisteix", "DataRegistre", "Gols", "Groga", "Vermella"]
        for c_idx, h in enumerate(headers_ass, 1):
            if ws_ass.cell(row=1, column=c_idx).value is None or str(ws_ass.cell(row=1, column=c_idx).value).strip() == "":
                ws_ass.cell(row=1, column=c_idx, value=h)

        def _inserir_registres(dict_ass, nom_equip):
            for jug_info, val_data in dict_ass.items():
                if isinstance(jug_info, tuple):
                    jug_info = dict(jug_info)
                nom_j = jug_info.get("nom", "") if isinstance(jug_info, dict) else ""
                id_p = jug_info.get("id_participant", "") if isinstance(jug_info, dict) else ""
                
                if isinstance(val_data, dict):
                    assist = val_data.get("assisteix", True)
                    gols = val_data.get("gols", 0)
                    groga = val_data.get("groga", 0)
                    vermella = val_data.get("vermella", 0)
                else:
                    assist = bool(val_data)
                    gols = 0
                    groga = 0
                    vermella = 0

                r_next = ws_ass.max_row + 1
                ws_ass.cell(row=r_next, column=1, value=str(id_partit))
                ws_ass.cell(row=r_next, column=2, value=str(id_p))
                ws_ass.cell(row=r_next, column=3, value=str(nom_j))
                ws_ass.cell(row=r_next, column=4, value=str(esport))
                ws_ass.cell(row=r_next, column=5, value=str(nom_equip))
                ws_ass.cell(row=r_next, column=6, value=1 if assist else 0)
                ws_ass.cell(row=r_next, column=7, value=now_str)
                ws_ass.cell(row=r_next, column=8, value=int(gols) if gols else None)
                ws_ass.cell(row=r_next, column=9, value=int(groga) if groga else None)
                ws_ass.cell(row=r_next, column=10, value=int(vermella) if vermella else None)

        _inserir_registres(assistencies_loc, eq_loc)
        _inserir_registres(assistencies_vis, eq_vis)
            
        wb.save(excel_file)
        return True, "Acta de partit i registre d'assistències desats amb èxit!"
    except Exception as e:
        return False, f"Error en desar a Excel: {e}"

# ==========================================
# 4. GENERACIÓ DE INFORMES (EXCEL & PDF)
# ==========================================

def generar_excel_resum(esport, grup, df_standings, df_pitxitxis, df_fp):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    font_title = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    font_sub = Font(name="Arial", size=9, italic=True, color="555555")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Arial", size=10)
    
    fill_title = PatternFill(start_color="008037", end_color="008037", fill_type="solid")
    fill_header_green = PatternFill(start_color="008037", end_color="008037", fill_type="solid")
    fill_header_blue = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FBF9", end_color="F9FBF9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    now_str = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    
    def afegir_pestanya(ws_title, df, fill_hdr):
        ws = wb.create_sheet(title=ws_title)
        num_cols = max(len(df.columns) if not df.empty else 4, 4)
        
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        cell_t = ws.cell(row=1, column=1, value=f"Resum Executiu Oficial — {esport} ({grup})")
        cell_t.font = font_title
        cell_t.fill = fill_title
        cell_t.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions.height = 28
        
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
        cell_s = ws.cell(row=2, column=1, value=f"SAF Servei d'Activitat Física UAB | Generat el {now_str}")
        cell_s.font = font_sub
        cell_s.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions.height = 18
        
        if df.empty:
            ws.cell(row=4, column=1, value="No hi ha dades registrades per a aquesta secció.").font = font_body
            return
            
        headers = list(df.columns)
        ws.row_dimensions.height = 22
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=4, column=col_idx, value=str(h))
            c.font = font_header
            c.fill = fill_hdr
            c.alignment = Alignment(horizontal='center', vertical='center')
            
        for row_idx, r in enumerate(df.itertuples(index=False), 5):
            ws.row_dimensions[row_idx].height = 18
            for col_idx, val in enumerate(r, 1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.font = font_body
                c.border = thin_border
                if row_idx % 2 == 0:
                    c.fill = fill_zebra
                c.alignment = Alignment(horizontal='left' if col_idx == 2 else 'center', vertical='center')
                
        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(headers[col_idx - 1]))
            for row_idx in range(5, 5 + len(df)):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

    afegir_pestanya('Classificacio', df_standings, fill_header_green)
    afegir_pestanya('Pitxitxis', df_pitxitxis, fill_header_blue)
    afegir_pestanya('FairPlay', df_fp, fill_header_green)

    wb.save(buffer)
    return buffer.getvalue()

def generar_pdf_resum(esport, grup, df_standings, df_pitxitxis, df_fp):
    esport = sanitize_pdf_text(esport)
    grup = sanitize_pdf_text(grup)
    
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'UABTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#008037'),
            alignment=1,
            spaceAfter=4
        )
        subtitle_style = ParagraphStyle(
            'UABSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            textColor=colors.HexColor('#555555'),
            alignment=1,
            spaceAfter=15
        )
        section_style = ParagraphStyle(
            'UABSection',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            textColor=colors.HexColor('#003366'),
            spaceBefore=10,
            spaceAfter=6
        )
        info_style = ParagraphStyle(
            'UABInfo',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            textColor=colors.HexColor('#777777'),
            spaceBefore=4,
            spaceAfter=6
        )
        
        story = []
        now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
        story.append(Paragraph(f"Resum Executiu Oficial - {esport} ({grup})", title_style))
        story.append(Paragraph(f"Generat el {now_str} | SAF Servei d'Activitat Fisica UAB", subtitle_style))
        
        # 1. Classificació
        story.append(Paragraph("1. Classificacio Oficial de Lliga", section_style))
        if not df_standings.empty:
            headers = ["Pos", "Equip", "Pts", "PJ", "PG", "PE", "PP", "GF", "GC", "DG"]
            data = [headers]
            for _, r in df_standings.iterrows():
                data.append([
                    sanitize_pdf_text(r["Pos"]), sanitize_pdf_text(r["Equip"])[:22], sanitize_pdf_text(r["Pts"]), sanitize_pdf_text(r["PJ"]),
                    sanitize_pdf_text(r["PG"]), sanitize_pdf_text(r["PE"]), sanitize_pdf_text(r["PP"]), sanitize_pdf_text(r["GF"]), sanitize_pdf_text(r["GC"]), sanitize_pdf_text(r["DG"])
                ])
            t = Table(data, colWidths=[30, 160, 35, 30, 30, 30, 30, 30, 30, 35])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#008037')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('ALIGN', (1,1), (1,-1), 'LEFT'),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('TOPPADDING', (0,0), (-1,-1), 4)
            ]))
            story.append(t)
        else:
            story.append(Paragraph("No hi ha equips o partits registrats per a aquesta seleccio.", info_style))
        story.append(Spacer(1, 10))
        
        # 2. Pitxitxis
        story.append(Paragraph("2. Ranquing de Pitxitxis (Golejadors)", section_style))
        if not df_pitxitxis.empty:
            data = [["Pos", "Jugador / Jugadora", "Gols"]]
            for idx, r in df_pitxitxis.head(10).iterrows():
                data.append([sanitize_pdf_text(idx), sanitize_pdf_text(r["Jugador"])[:35], sanitize_pdf_text(r["Gols"])])
            t = Table(data, colWidths=[40, 300, 60])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#003366')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('ALIGN', (1,1), (1,-1), 'LEFT'),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4)
            ]))
            story.append(t)
        else:
            story.append(Paragraph("No s'ha registrat cap golejador en els partits validats d'aquesta competicio.", info_style))
        story.append(Spacer(1, 10))
        
        # 3. Fair Play
        story.append(Paragraph("3. Classificacio Fair Play (Jugar Net)", section_style))
        if not df_fp.empty:
            data = [["Pos", "Equip", "PJ", "Grogues", "Vermelles", "Pts Sancio"]]
            for idx, r in df_fp.iterrows():
                data.append([sanitize_pdf_text(idx), sanitize_pdf_text(r["Equip"])[:25], sanitize_pdf_text(r["PJ"]), sanitize_pdf_text(r["Grogues"]), sanitize_pdf_text(r["Vermelles"]), sanitize_pdf_text(r["Punts Sanció"])])
            t = Table(data, colWidths=[40, 200, 50, 60, 60, 70])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#008037')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('ALIGN', (1,1), (1,-1), 'LEFT'),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4)
            ]))
            story.append(t)
        else:
            story.append(Paragraph("No hi ha dades de targetes o Fair Play per a aquesta competicio.", info_style))
            
        doc.build(story)
        return buffer.getvalue()
    except Exception as e_rl:
        try:
            from fpdf import FPDF
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 14)
            pdf.cell(0, 10, sanitize_pdf_text(f"Resum Executiu - {esport} ({grup})"), 0, 1, "C")
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 8, "Classificacio Oficial", 0, 1)
            for _, r in df_standings.iterrows():
                txt = sanitize_pdf_text(f"{r['Pos']}. {r['Equip']} - Pts: {r['Pts']} (PJ:{r['PJ']} GF:{r['GF']} GC:{r['GC']})")
                pdf.cell(0, 6, txt, 0, 1)
            
            res_out = pdf.output(dest='S')
            if isinstance(res_out, str):
                return res_out.encode('latin1')
            return bytes(res_out)
        except Exception as e_fpdf:
            raise RuntimeError(f"No s'ha pogut generar el PDF: {e_rl} | Fallback: {e_fpdf}")

# ==========================================
# 5. EXECUCIÓ I NAVEGACIÓ PER PESTANYES
# ==========================================

df_cal = carregar_calendari()

df_standings = pd.DataFrame()
df_pitxitxis = pd.DataFrame()
df_fp = pd.DataFrame()
esport_sel = ""
grup_sel = ""

tab_public, tab_pitxitxi, tab_fairplay, tab_arbitre = st.tabs([
    "📊 Classificació Lliga",
    "⚽ Pitxitxis (Golejadors)",
    "🕊️ Fair Play (Jugar Net)",
    "🔒 Àrea Privada Àrbitres"
])

# ==========================================
# PESTANYA 1: PORTAL PÚBLIC (CLASSIFICACIÓ)
# ==========================================
with tab_public:
    if not df_cal.empty and "Esport" in df_cal.columns and "Grup" in df_cal.columns:
        esports = df_cal["Esport"].dropna().unique()
        
        if len(esports) > 0:
            col1, col2 = st.columns(2)
            with col1:
                esport_sel = st.selectbox("🏆 Selecciona Competició:", esports, key="pub_esp")
            with col2:
                df_esp = df_cal[df_cal["Esport"] == esport_sel]
                grups = df_esp["Grup"].dropna().unique()
                grup_sel = st.selectbox("👥 Selecciona Grup:", grups, key="pub_grp")
                
            df_filtrat = df_esp[df_esp["Grup"] == grup_sel]
            
            teams = list(set(df_filtrat["EquipLocal"].dropna().tolist() + df_filtrat["EquipVisitant"].dropna().tolist()))
            teams = [t for t in teams if t != "DESCANSA"]
            
            standings = []
            for t in teams:
                df_home = df_filtrat[(df_filtrat["EquipLocal"] == t) & (df_filtrat["EstatPartit"].astype(str).str.strip().str.upper().isin(["JUGAT", "VALIDAT"]))]
                df_away = df_filtrat[(df_filtrat["EquipVisitant"] == t) & (df_filtrat["EstatPartit"].astype(str).str.strip().str.upper().isin(["JUGAT", "VALIDAT"]))]
                
                pj = len(df_home) + len(df_away)
                pg, pe, pp, gf, gc = 0, 0, 0, 0, 0
                
                for _, row in df_home.iterrows():
                    rl = safe_int(row.get("ResultatLocal"))
                    rv = safe_int(row.get("ResultatVisitant"))
                    gf += rl; gc += rv
                    if rl > rv: pg += 1
                    elif rl == rv: pe += 1
                    else: pp += 1
                    
                for _, row in df_away.iterrows():
                    rl = safe_int(row.get("ResultatLocal"))
                    rv = safe_int(row.get("ResultatVisitant"))
                    gf += rv; gc += rl
                    if rv > rl: pg += 1
                    elif rv == rl: pe += 1
                    else: pp += 1
                    
                pts = pg * 3 + pe * 1
                dg = gf - gc
                
                standings.append({
                    "Pos": 0, "Equip": t, "Pts": pts, "PJ": pj,
                    "PG": pg, "PE": pe, "PP": pp, "GF": gf, "GC": gc, "DG": dg
                })
                
            df_standings = pd.DataFrame(standings)
            if not df_standings.empty:
                df_standings = df_standings.sort_values(by=["Pts", "DG", "GF", "Equip"], ascending=[False, False, False, True])
                df_standings["Pos"] = range(1, len(df_standings) + 1)
                
            # Pitxitxis
            dict_goals = {}
            df_validats = df_filtrat[df_filtrat["EstatPartit"].astype(str).str.strip().str.upper().isin(["JUGAT", "VALIDAT"])]
            
            col_gl_name = None
            col_gv_name = None
            for c in df_validats.columns:
                if "golejador" in c.lower() and "local" in c.lower(): col_gl_name = c
                elif "golejador" in c.lower() and "visitant" in c.lower(): col_gv_name = c
                
            if not col_gl_name and len(df_validats.columns) >= 22: col_gl_name = df_validats.columns[21]
            if not col_gv_name and len(df_validats.columns) >= 23: col_gv_name = df_validats.columns[22]
            
            for _, row in df_validats.iterrows():
                val_loc = str(row.get(col_gl_name, "")) if col_gl_name else ""
                val_vis = str(row.get(col_gv_name, "")) if col_gv_name else ""
                
                for text_gols in [val_loc, val_vis]:
                    if text_gols and text_gols.lower() != "nan" and text_gols.lower() != "none":
                        for part in text_gols.split(","):
                            part = part.strip()
                            if part:
                                jugador = part
                                num_gols = 1
                                if "(" in part and ")" in part:
                                    try:
                                        jugador = part[:part.find("(")].strip()
                                        content = part[part.find("(")+1:part.rfind(")")].strip()
                                        match = re.search(r'\d+', content)
                                        if match:
                                            num_gols = int(match.group())
                                    except Exception:
                                        pass
                                dict_goals[jugador] = dict_goals.get(jugador, 0) + num_gols
                                
            df_pitxitxis = pd.DataFrame(list(dict_goals.items()), columns=["Jugador", "Gols"])
            if not df_pitxitxis.empty:
                df_pitxitxis = df_pitxitxis.sort_values(by=["Gols", "Jugador"], ascending=[False, True])
                df_pitxitxis.index = range(1, len(df_pitxitxis) + 1)
                
            # Fair Play
            dict_fp = {t: {"PJ": 0, "Grogues": 0, "Vermelles": 0, "Punts Sanció": 0} for t in teams}
            for _, row in df_validats.iterrows():
                eq_l, eq_v = row.get("EquipLocal"), row.get("EquipVisitant")
                tgl, tvl = safe_int(row.get("TargetesGroguesLocal")), safe_int(row.get("TargetesVermellesLocal"))
                tgv, tvv = safe_int(row.get("TargetesGroguesVisitant")), safe_int(row.get("TargetesVermellesVisitant"))
                
                if eq_l in dict_fp:
                    dict_fp[eq_l]["PJ"] += 1
                    dict_fp[eq_l]["Grogues"] += tgl
                    dict_fp[eq_l]["Vermelles"] += tvl
                    dict_fp[eq_l]["Punts Sanció"] += (tgl * 1 + tvl * 3)
                    
                if eq_v in dict_fp:
                    dict_fp[eq_v]["PJ"] += 1
                    dict_fp[eq_v]["Grogues"] += tgv
                    dict_fp[eq_v]["Vermelles"] += tvv
                    dict_fp[eq_v]["Punts Sanció"] += (tgv * 1 + tvv * 3)
                    
            df_fp = pd.DataFrame([{"Equip": k, **v} for k, v in dict_fp.items()])
            if not df_fp.empty:
                df_fp = df_fp.sort_values(by=["Punts Sanció", "Vermelles", "Grogues", "Equip"], ascending=[True, True, True, True])
                df_fp.index = range(1, len(df_fp) + 1)

            st.markdown("---")
            c_pdf1, c_pdf2, c_pdf3 = st.columns([2, 1, 1])
            with c_pdf1:
                st.markdown(f"##### 📥 Baixar Resums Executius - {esport_sel} ({grup_sel})")
            
            with c_pdf2:
                try:
                    excel_bytes = generar_excel_resum(esport_sel, grup_sel, df_standings, df_pitxitxis, df_fp)
                    st.download_button(
                        label="📊 Resum Excel",
                        data=excel_bytes,
                        file_name=f"Resum_Executiu_{esport_sel}_{grup_sel}.xlsx".replace(" ", "_"),
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
                except Exception as e_ex:
                    st.error(f"Error Excel: {e_ex}")
                    
            with c_pdf3:
                try:
                    pdf_bytes = generar_pdf_resum(esport_sel, grup_sel, df_standings, df_pitxitxis, df_fp)
                    st.download_button(
                        label="📄 Resum PDF",
                        data=pdf_bytes,
                        file_name=f"Resum_Executiu_{esport_sel}_{grup_sel}.pdf".replace(" ", "_"),
                        mime="application/pdf"
                    )
                except Exception as e_pdf:
                    st.warning("💡 Solució PDF: Executa `pip install reportlab` a la teva terminal.")
            st.markdown("---")

            if not df_standings.empty:
                st.dataframe(df_standings.set_index("Pos"), use_container_width=True)
            else:
                st.info("No hi ha equips o partits registrats per a aquesta selecció.")
                
            st.markdown("#### 📅 Calendari de Partits & Resultats")
            cols_cal = [c for c in ["Id", "Jornada", "DataPartit", "HoraPartit", "Camp", "EquipLocal", "ResultatLocal", "ResultatVisitant", "EquipVisitant", "EstatPartit"] if c in df_filtrat.columns]
            st.dataframe(df_filtrat[cols_cal], use_container_width=True)
    else:
        st.info("ℹ️ Carregant dades del calendari. Assegura't de tenir el fitxer 'Plantilla_MultiEsports_UAB.xlsx' a la mateixa carpeta.")

# ==========================================
# PESTANYA 2: PITXITXIS
# ==========================================
with tab_pitxitxi:
    if esport_sel and grup_sel:
        st.markdown(f"### ⚽ Rànquing Oficial de Golejadors - {esport_sel} ({grup_sel})")
    else:
        st.markdown("### ⚽ Rànquing Oficial de Golejadors")
        
    if not df_pitxitxis.empty:
        st.dataframe(df_pitxitxis, use_container_width=True)
    else:
        st.info("No s'ha registrat cap golejador en els partits validats d'aquesta competició.")

# ==========================================
# PESTANYA 3: FAIR PLAY
# ==========================================
with tab_fairplay:
    if esport_sel and grup_sel:
        st.markdown(f"### 🕊️ Classificació Fair Play (Jugar Net) - {esport_sel} ({grup_sel})")
    else:
        st.markdown("### 🕊️ Classificació Fair Play (Jugar Net)")
        
    st.caption("💡 Càlcul de Punts de Sanció: Targeta Groga = 1 punt | Targeta Vermella = 3 punts (Guanya l'equip amb menys punts)")
    if not df_fp.empty:
        st.dataframe(df_fp, use_container_width=True)
    else:
        st.info("No hi ha dades de targetes o Fair Play per a aquesta competició.")

# ==========================================
# PESTANYA 4: ÀREA PRIVADA D'ÀRBITRES I ASSISTÈNCIES
# ==========================================
with tab_arbitre:
    st.markdown("### 🔐 Accés Restringit per a Àrbitres i Organització")
    
    pwd_input = st.text_input("Introdueix la contrasenya d'accés:", type="password", help="Contrasenya per defecte: UAB2026")
    
    if pwd_input == PASSWORD_ARBITRE:
        st.success("Sessió d'Àrbitre Validada amb Èxit! ⚽")
        
        if not df_cal.empty:
            df_pendents = df_cal[~df_cal["EstatPartit"].astype(str).str.strip().str.upper().isin(["VALIDAT"])]
            
            if not df_pendents.empty:
                st.markdown("#### 📝 Introducció d'Acta i Validació d'Assistències")
                
                partit_dict = {}
                for _, r in df_pendents.iterrows():
                    label = f"[{r['Id']}] {r['Esport']} ({r['Grup']}) | J{r['Jornada']} | {r['EquipLocal']} vs {r['EquipVisitant']} ({r['DataPartit']} {r['HoraPartit']})"
                    partit_dict[label] = r
                    
                sel_partit_label = st.selectbox("Selecciona el Partit a Validar:", list(partit_dict.keys()))
                row_sel = partit_dict[sel_partit_label]
                
                p_id = row_sel["Id"]
                p_esport = row_sel["Esport"]
                p_eq_loc = row_sel["EquipLocal"]
                p_eq_vis = row_sel["EquipVisitant"]
                
                st.info(f"**Competició:** {p_esport} — **Grup:** {row_sel['Grup']} (Jornada {row_sel['Jornada']}) | **{p_eq_loc}** vs **{p_eq_vis}**")
                
                # Càrrega de jugadors de Consolidat
                jugadors_loc = carregar_jugadors_equip(p_esport, p_eq_loc)
                jugadors_vis = carregar_jugadors_equip(p_esport, p_eq_vis)
                
                with st.form("form_acta_completa"):
                    st.markdown("#### 1. Resultat i Marcador Oficial")
                    c1, c2 = st.columns(2)
                    with c1:
                        st.markdown(f"##### 🏠 {p_eq_loc} (LOCAL)")
                        gols_loc = st.number_input("Gols Local:", min_value=0, max_value=99, value=safe_int(row_sel.get("ResultatLocal")), step=1)
                        t_grogues_loc = st.number_input("Targetes Grogues Local:", min_value=0, max_value=20, value=0, step=1)
                        t_vermelles_loc = st.number_input("Targetes Vermelles Local:", min_value=0, max_value=10, value=0, step=1)
                        golejadors_loc = st.text_input("Golejadors Local (ex: Pérez (2), Marc (1)):", value="")
                        
                    with c2:
                        st.markdown(f"##### ✈️ {p_eq_vis} (VISITANT)")
                        gols_vis = st.number_input("Gols Visitant:", min_value=0, max_value=99, value=safe_int(row_sel.get("ResultatVisitant")), step=1)
                        t_grogues_vis = st.number_input("Targetes Grogues Visitant:", min_value=0, max_value=20, value=0, step=1)
                        t_vermelles_vis = st.number_input("Targetes Vermelles Visitant:", min_value=0, max_value=10, value=0, step=1)
                        golejadors_vis = st.text_input("Golejadors Visitant (ex: Joan (1)):", value="")
                        
                    st.markdown("---")
                    st.markdown("#### 2. Control d'Assistència de Jugadors/es (Full Consolidat)")
                    
                    ca_loc, ca_vis = st.columns(2)
                    dict_ass_loc = {}
                    dict_ass_vis = {}
                    
                    with ca_loc:
                        st.markdown(f"##### 📋 Convocatòria & Estadístiques {p_eq_loc}")
                        if jugadors_loc:
                            for idx_j, jug in enumerate(jugadors_loc):
                                c_j1, c_j2, c_j3, c_j4 = st.columns([3, 1, 1, 1])
                                with c_j1:
                                    chk = st.checkbox(f"{jug['nom']}", value=True, key=f"chk_loc_{p_id}_{idx_j}")
                                with c_j2:
                                    gols_j = st.number_input("Gols", min_value=0, max_value=20, value=0, key=f"gols_loc_{p_id}_{idx_j}")
                                with c_j3:
                                    groga_j = st.number_input("🟨", min_value=0, max_value=2, value=0, key=f"groga_loc_{p_id}_{idx_j}")
                                with c_j4:
                                    verm_j = st.number_input("🟥", min_value=0, max_value=1, value=0, key=f"verm_loc_{p_id}_{idx_j}")
                                dict_ass_loc[tuple(jug.items())] = {"assisteix": chk, "gols": gols_j, "groga": groga_j, "vermella": verm_j}
                        else:
                            st.caption("No s'ha trobat llista de jugadors per a aquest equip al full Consolidat.")
                            
                    with ca_vis:
                        st.markdown(f"##### 📋 Convocatòria & Estadístiques {p_eq_vis}")
                        if jugadors_vis:
                            for idx_j, jug in enumerate(jugadors_vis):
                                c_j1, c_j2, c_j3, c_j4 = st.columns([3, 1, 1, 1])
                                with c_j1:
                                    chk = st.checkbox(f"{jug['nom']}", value=True, key=f"chk_vis_{p_id}_{idx_j}")
                                with c_j2:
                                    gols_j = st.number_input("Gols", min_value=0, max_value=20, value=0, key=f"gols_vis_{p_id}_{idx_j}")
                                with c_j3:
                                    groga_j = st.number_input("🟨", min_value=0, max_value=2, value=0, key=f"groga_vis_{p_id}_{idx_j}")
                                with c_j4:
                                    verm_j = st.number_input("🟥", min_value=0, max_value=1, value=0, key=f"verm_vis_{p_id}_{idx_j}")
                                dict_ass_vis[tuple(jug.items())] = {"assisteix": chk, "gols": gols_j, "groga": groga_j, "vermella": verm_j}
                        else:
                            st.caption("No s'ha trobat llista de jugadors per a aquest equip al full Consolidat.")
                            
                    st.markdown("---")
                    submitted = st.form_submit_button("💾 Validar Acta, Desar Assistències i Actualitzar Classificació")
                    
                    if submitted:
                        # Reconvertir claus de diccionari de tuples a dicts
                        ass_loc_clean = {dict(k): v for k, v in dict_ass_loc.items()}
                        ass_vis_clean = {dict(k): v for k, v in dict_ass_vis.items()}
                        
                        ok, msg = guardar_acta_amb_assistencies(
                            p_id, p_esport, p_eq_loc, p_eq_vis,
                            gols_loc, gols_vis,
                            t_grogues_loc, t_vermelles_loc, t_grogues_vis, t_vermelles_vis,
                            golejadors_loc, golejadors_vis,
                            ass_loc_clean, ass_vis_clean
                        )
                        if ok:
                            st.success(msg)
                            st.rerun()
                        else:
                            st.error(msg)
            else:
                st.info("Tots els partits programats ja estan validats!")
    elif pwd_input != "":
        st.error("Contrasenya incorrecta. Revisa les teves credencials d'àrbitre.")
